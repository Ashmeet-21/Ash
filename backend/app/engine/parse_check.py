"""ATS parse check: what a basic resume parser actually extracts from your file.

Checks the known parse breakers: text that can't be extracted, multi-column layouts,
tables, images, text boxes, and contact details placed in headers or footers.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import asdict, dataclass, field

import pdfplumber
from docx import Document

from app.engine.validator import validate_text_rules

EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE = re.compile(r"(?:\+?\d{1,3}[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}")
STANDARD_HEADINGS = ["experience", "projects", "education", "skills"]
HEADER_BAND = 0.06  # top/bottom fraction of the page where Word headers and footers render


@dataclass
class Check:
    key: str
    passed: bool
    severity: str  # error | warn | info
    message: str


@dataclass
class ParseResult:
    file_type: str
    pages: int
    text: str
    checks: list[Check] = field(default_factory=list)
    detected_sections: list[str] = field(default_factory=list)
    contact: dict = field(default_factory=dict)

    @property
    def format_score(self) -> float:
        penalty = {"error": 0.35, "warn": 0.12, "info": 0.0}
        return round(max(0.0, 1.0 - sum(penalty[c.severity] for c in self.checks if not c.passed)), 2)

    def to_dict(self) -> dict:
        d = asdict(self)
        d["format_score"] = self.format_score
        return d


def _common_checks(res: ParseResult) -> None:
    text = res.text
    lower = text.lower()
    res.detected_sections = [h for h in STANDARD_HEADINGS if re.search(rf"(?m)^\s*{h}\b", lower)]
    missing = [h for h in STANDARD_HEADINGS if h not in res.detected_sections]
    res.checks.append(Check("standard_headings", not missing, "warn", "Standard headings found." if not missing else f"Missing standard headings: {', '.join(missing)}."))
    res.contact = {"email": (EMAIL.search(text) or [None])[0], "phone": (PHONE.search(text) or [None])[0]}
    res.checks.append(Check("contact_found", bool(res.contact["email"]), "error", "Email extracted." if res.contact["email"] else "No email address could be extracted."))
    broken = text.count("(cid:") + text.count("�")
    res.checks.append(Check("encoding", broken == 0, "error", "Characters decode cleanly." if not broken else f"{broken} characters failed to decode (font encoding issue)."))
    rules = validate_text_rules(text)
    for issue in rules.errors:
        res.checks.append(Check(issue.rule, False, "warn", issue.message))


def _columns_on_page(page) -> float:
    """Fraction of text lines that have a wide gap crossing the page's middle (a two-column tell)."""
    words = page.extract_words(keep_blank_chars=False, use_text_flow=False)
    if not words:
        return 0.0
    width = float(page.width)
    lines: dict[int, list] = {}
    for w in words:
        lines.setdefault(round(float(w["top"]) / 3), []).append(w)
    split = 0
    for ws in lines.values():
        ws.sort(key=lambda w: w["x0"])
        for a, b in zip(ws, ws[1:]):
            gap = float(b["x0"]) - float(a["x1"])
            if gap > 0.08 * width and float(a["x1"]) < 0.62 * width and float(b["x0"]) > 0.3 * width:
                split += 1
                break
    return split / len(lines)


def check_pdf(data: bytes) -> ParseResult:
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = pdf.pages
        text = "\n".join((p.extract_text() or "") for p in pages)
        res = ParseResult("pdf", len(pages), text)
        chars = len(text.strip())
        res.checks.append(Check("text_extractable", chars > 300, "error", f"{chars} characters extracted." if chars > 300 else "Little or no text extracted; the PDF may be scanned or image-based."))

        col_frac = max((_columns_on_page(p) for p in pages), default=0.0)
        res.checks.append(Check("single_column", col_frac < 0.25, "error" if col_frac >= 0.4 else "warn",
                                "Single-column layout." if col_frac < 0.25 else f"{col_frac:.0%} of lines look split into columns; parsers may interleave them."))

        n_tables = sum(len(p.find_tables()) for p in pages)
        res.checks.append(Check("no_tables", n_tables == 0, "warn", "No tables." if not n_tables else f"{n_tables} table(s) found; parsers often scramble table cells."))
        n_images = sum(len(p.images) for p in pages)
        res.checks.append(Check("no_images", n_images == 0, "warn", "No images." if not n_images else f"{n_images} image(s); any text inside them is invisible to an ATS."))

        # Contact details only in the header/footer band.
        in_band = in_body = False
        for p in pages:
            h = float(p.height)
            for w in p.extract_words():
                if EMAIL.search(w["text"]) or PHONE.search(w["text"]):
                    if float(w["top"]) < HEADER_BAND * h or float(w["bottom"]) > (1 - HEADER_BAND) * h:
                        in_band = True
                    else:
                        in_body = True
        ok = in_body or not in_band
        res.checks.append(Check("contact_in_body", ok, "warn", "Contact details sit in the document body." if ok else "Contact details appear only in the page header/footer area, which some parsers skip."))
        if len(pages) > 2:
            res.checks.append(Check("length", False, "warn", f"{len(pages)} pages; new-grad resumes should be 1 page."))
    _common_checks(res)
    return res


def check_docx(data: bytes) -> ParseResult:
    doc = Document(io.BytesIO(data))
    body = [p.text for p in doc.paragraphs]
    res = ParseResult("docx", 0, "\n".join(body))
    header_text = " ".join(p.text for s in doc.sections for part in (s.header, s.footer) for p in part.paragraphs)
    res.checks.append(Check("no_tables", not doc.tables, "warn", "No tables." if not doc.tables else f"{len(doc.tables)} table(s); parsers often scramble table cells."))
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read("word/document.xml").decode("utf8", "ignore")
    n_boxes = xml.count("<w:txbxContent")
    res.checks.append(Check("no_text_boxes", n_boxes == 0, "error", "No text boxes." if not n_boxes else f"{n_boxes} text box(es); their text is often dropped."))
    cols = re.findall(r'<w:cols[^>]*w:num="(\d+)"', xml)
    multi = any(int(c) > 1 for c in cols)
    res.checks.append(Check("single_column", not multi, "error", "Single-column layout." if not multi else "Multi-column section layout."))
    n_images = len(doc.inline_shapes)
    res.checks.append(Check("no_images", n_images == 0, "warn", "No images." if not n_images else f"{n_images} image(s)."))
    contact_in_header = bool(EMAIL.search(header_text) or PHONE.search(header_text))
    contact_in_body = bool(EMAIL.search(res.text) or PHONE.search(res.text))
    ok = contact_in_body or not contact_in_header
    res.checks.append(Check("contact_in_body", ok, "warn", "Contact details sit in the document body." if ok else "Contact details are only in the Word header/footer."))
    _common_checks(res)
    return res


def parse_check(filename: str, data: bytes) -> ParseResult:
    name = filename.lower()
    if name.endswith(".pdf") or data[:4] == b"%PDF":
        return check_pdf(data)
    if name.endswith(".docx"):
        return check_docx(data)
    raise ValueError("Upload a .pdf or .docx file.")

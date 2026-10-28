"""Builds the tailored resume as an ATS-safe DOCX (single column, standard headings,
contact in the body, no tables or text boxes), converts it to PDF with LibreOffice when
available, and re-validates the text extracted from the output file."""

from __future__ import annotations

import io
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt

from app.engine.parse_check import parse_check
from app.engine.validator import ValidationResult, validate_document
from app.facts import Facts, LockedEntry

FONT = "Calibri"


def _para(doc, text: str = "", bold: bool = False, size: float = 10.5, space_after: float = 0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        run.font.name = FONT
    return p


def _heading(doc, text: str):
    p = _para(doc, text.upper(), bold=True, size=11.5)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)


def _entry(doc, e: LockedEntry, right_tab: float):
    p = _para(doc)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(right_tab), WD_TAB_ALIGNMENT.RIGHT)
    head = p.add_run(e.title)
    head.bold = True
    head.font.size = Pt(10.5)
    rest = " | ".join(x for x in [e.org, e.tech] if x)
    if rest:
        r = p.add_run(f" | {rest}")
        r.font.size = Pt(10.5)
    if e.dates:
        r = p.add_run(f"\t{e.dates}")
        r.font.size = Pt(10.5)
    for b in e.bullets:
        bp = doc.add_paragraph(style="List Bullet")
        bp.paragraph_format.space_after = Pt(0)
        run = bp.add_run(b)
        run.font.size = Pt(10.5)


def build_docx(facts: Facts, objective: str, skills: dict[str, list[str]]) -> bytes:
    doc = Document()
    section = doc.sections[0]
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Inches(0.6))
    section.top_margin = section.bottom_margin = Inches(0.5)
    right_tab = (section.page_width - section.left_margin - section.right_margin) / 914400  # EMU per inch
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10.5)

    p, r = facts.profile, facts.resume
    for name in r.section_order:
        key = name.lower()
        if key == "header":
            _para(doc, p.name, bold=True, size=16)
            contact = " | ".join(x for x in [p.contact.email, p.contact.phone, p.contact.linkedin, p.contact.github, p.contact.location] if x)
            _para(doc, contact, size=10, space_after=4)
        elif key == "objective":
            _heading(doc, name)
            _para(doc, " ".join(objective.split()))
        elif key == "skills":
            _heading(doc, name)
            for label, items in skills.items():
                para = _para(doc)
                lr = para.add_run(f"{label}: ")
                lr.bold = True
                para.add_run(", ".join(items))
        elif key in ("projects", "experience", "education"):
            _heading(doc, name)
            for e in getattr(r, key):
                _entry(doc, e, right_tab)
        elif key == "certifications":
            _heading(doc, name)
            for c in r.certifications:
                _para(doc, c)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def docx_text(data: bytes) -> str:
    return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)


def to_pdf(docx_bytes: bytes) -> bytes | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "resume.docx"
        src.write_bytes(docx_bytes)
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)],
                check=True, capture_output=True, timeout=120,
            )
        except (subprocess.SubprocessError, OSError):
            return None
        out = Path(tmp) / "resume.pdf"
        return out.read_bytes() if out.exists() else None


def build_and_check(facts: Facts, objective: str, skills: dict[str, list[str]], out_dir: Path, stem: str) -> dict:
    """Write DOCX (+PDF) and validate what a parser reads back from them."""
    out_dir.mkdir(parents=True, exist_ok=True)
    docx_bytes = build_docx(facts, objective, skills)
    docx_path = out_dir / f"{stem}.docx"
    docx_path.write_bytes(docx_bytes)
    validation: ValidationResult = validate_document(docx_text(docx_bytes), facts)

    result: dict = {"docx_path": str(docx_path), "pdf_path": None, "parse_check": None}
    pdf = to_pdf(docx_bytes)
    if pdf:
        pdf_path = out_dir / f"{stem}.pdf"
        pdf_path.write_bytes(pdf)
        result["pdf_path"] = str(pdf_path)
        pc = parse_check(pdf_path.name, pdf)
        result["parse_check"] = pc.to_dict()
        validation.merge(validate_document(pc.text, facts))
    result["validation"] = validation.to_dict()
    return result
